"""
Shared column contract and workbook styling.

Every script agrees on these names. If a real export uses different headers, map them
once here (or pass --column-map) rather than editing four scripts and hoping they stay
in step.
"""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# What the classifier reads. Deliberately narrow, see references/classification-rules.md.
TEXT_COLUMNS = ["Short Description", "Description"]
CONTEXT_COLUMNS = ["Department", "Channel"]

# What the analyser needs. Missing ones degrade specific sections rather than crashing.
REQUIRED = ["Ticket ID", "Month", "Description"]
OPTIONAL = ["Opened Date", "Short Description", "Requester", "Department", "Channel",
            "Priority", "Assignment Group", "State", "Resolved Date", "Resolution Hours",
            "SLA Target Hours", "SLA Met", "Reopened"]
CATEGORY_COLUMNS = ["Category (L1)", "Subcategory (L2)", "Issue Type (L3)"]

PRIORITIES = ["P1 - Critical", "P2 - High", "P3 - Moderate", "P4 - Low"]
SLA_TARGET = {"P1 - Critical": 4, "P2 - High": 8, "P3 - Moderate": 16, "P4 - Low": 40}
CHANNELS = ["Self-Service Portal", "Email", "Phone", "Chat"]

WIDTHS = {"Description": 62, "Short Description": 30, "Reason": 52, "Ticket ID": 12,
          "Opened Date": 17, "Resolved Date": 17, "Requester": 18, "Priority": 15,
          "Department": 18, "Channel": 19, "Assignment Group": 20,
          "Category (L1)": 24, "Subcategory (L2)": 26, "Issue Type (L3)": 30,
          "True Category (L1)": 24, "True Subcategory (L2)": 26,
          "True Issue Type (L3)": 30, "Example Phrasing": 55}


def check_columns(df, need_categories=False):
    """Report what is missing and what that costs, instead of failing on a KeyError
    three scripts later."""
    missing = [c for c in REQUIRED if c not in df.columns]
    if need_categories:
        missing += [c for c in CATEGORY_COLUMNS if c not in df.columns]
    if missing:
        raise SystemExit(
            "The export is missing required columns: " + ", ".join(missing) +
            "\nRename them in the source file, or pass --column-map "
            "'Their Name=Our Name,...' to map them.")
    degraded = [c for c in OPTIONAL if c not in df.columns]
    return degraded


def apply_column_map(df, spec):
    """spec looks like 'number=Ticket ID,opened_at=Opened Date'."""
    if not spec:
        return df
    pairs = [p.split("=", 1) for p in spec.split(",") if "=" in p]
    return df.rename(columns={a.strip(): b.strip() for a, b in pairs})


def style_workbook(book, header="1F3B57"):
    fill = PatternFill("solid", fgColor=header)
    font = Font(color="FFFFFF", bold=True, size=10)
    for ws in book.worksheets:
        if ws.max_row < 1:
            continue
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(vertical="center")
        for i, cell in enumerate(ws[1], start=1):
            ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(
                cell.value, max(14, min(30, len(str(cell.value)) + 4)))
        ws.auto_filter.ref = ws.dimensions
