"""
Step 2c of the workflow: fan the case-level classifications back out to every ticket,
validate them against the taxonomy, and score them against the held-back Answer Key.

The scoring step is the point. Without it we would be reporting a categorised dataset
with no idea whether the categories are right, and every number downstream would be
an assertion rather than a measurement.

Output:
  data/categorized-tickets.xlsx - Tickets sheet with the three category columns filled,
                                  plus an Accuracy sheet and a Needs Review sheet
  data/accuracy.json            - the measured accuracy, consumed by the dashboard

Run: python3 apply_classification.py
"""

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from taxonomy import allowed_paths

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
SRC = DATA / "servicenow-tickets-jan-mar-2026.xlsx"


def main():
    cases = pd.read_json(DATA / "cases.json")
    tmap = pd.read_json(DATA / "ticket_case_map.json")
    tickets = pd.read_excel(SRC, sheet_name="Tickets")
    key = pd.read_excel(SRC, sheet_name="Answer Key").drop(columns=["Month"])

    preds = [json.loads(line) for line in
             (DATA / "classified.jsonl").read_text().splitlines() if line.strip()]
    preds = pd.DataFrame(preds)

    # Gate 1: every case must have been classified, exactly once.
    missing = set(cases.case_id) - set(preds.case_id)
    extra = set(preds.case_id) - set(cases.case_id)
    if missing or extra:
        raise SystemExit(f"Case mismatch. Missing: {sorted(missing)} Unknown: {sorted(extra)}")
    if preds.case_id.duplicated().any():
        raise SystemExit("Duplicate case_id in classified.jsonl")

    # Gate 2: every path must exist in the taxonomy, character for character.
    valid = set(allowed_paths()) | {"UNCATEGORIZED"}
    bad = preds[~preds.path.isin(valid)]
    if len(bad):
        raise SystemExit(f"Paths not in taxonomy:\n{bad[['case_id', 'path']].to_string()}")

    parts = preds.path.str.split(" > ", expand=True)
    preds["Category (L1)"] = parts[0].fillna("Uncategorized")
    preds["Subcategory (L2)"] = parts[1].fillna("Uncategorized")
    preds["Issue Type (L3)"] = parts[2].fillna("Uncategorized")

    case_lookup = cases[["case_id", "_norm"]].merge(preds, on="case_id")
    out = (tickets.drop(columns=["Category (L1)", "Subcategory (L2)", "Issue Type (L3)"])
           .merge(tmap, on="Ticket ID")
           .merge(case_lookup, on="_norm", how="left")
           .drop(columns=["_norm"])
           .rename(columns={"confidence": "Confidence", "reason": "Reason",
                            "case_id": "Case ID"}))

    truth = out.merge(key, on="Ticket ID")
    for lvl, col in [("L1", "Category (L1)"), ("L2", "Subcategory (L2)"),
                     ("L3", "Issue Type (L3)")]:
        truth[f"correct_{lvl}"] = truth[col] == truth[f"True {col}"]

    n = len(truth)
    overall = {lvl: round(100 * truth[f"correct_{lvl}"].sum() / n, 1)
               for lvl in ["L1", "L2", "L3"]}

    per_cat = (truth.groupby("True Category (L1)")
               .agg(tickets=("Ticket ID", "count"),
                    L1_correct_pct=("correct_L1", lambda s: round(100 * s.mean(), 1)),
                    L2_correct_pct=("correct_L2", lambda s: round(100 * s.mean(), 1)),
                    L3_correct_pct=("correct_L3", lambda s: round(100 * s.mean(), 1)))
               .reset_index().rename(columns={"True Category (L1)": "Category (L1)"}))

    by_conf = (truth.groupby("Confidence")
               .agg(tickets=("Ticket ID", "count"),
                    L2_correct_pct=("correct_L2", lambda s: round(100 * s.mean(), 1)))
               .reindex(["High", "Medium", "Low"]).dropna(how="all").reset_index())

    review = out[out.Confidence.isin(["Low", "Medium"])][
        ["Ticket ID", "Month", "Short Description", "Description",
         "Category (L1)", "Subcategory (L2)", "Issue Type (L3)", "Confidence", "Reason"]]

    wrong = truth[~truth.correct_L2][
        ["Ticket ID", "Description", "Subcategory (L2)", "True Subcategory (L2)",
         "Confidence"]].drop_duplicates(subset=["Description"])

    accuracy = {
        "tickets_scored": int(n),
        "unique_cases": int(len(cases)),
        "classification_calls_saved_pct": round(100 * (1 - len(cases) / n), 1),
        "overall": overall,
        "per_category": per_cat.to_dict(orient="records"),
        "by_confidence": by_conf.to_dict(orient="records"),
        "uncategorized": int((out["Category (L1)"] == "Uncategorized").sum()),
        "needs_review_tickets": int(len(review)),
        "distinct_misclassified_cases": int(len(wrong)),
    }
    (DATA / "accuracy.json").write_text(json.dumps(accuracy, indent=2))

    path = DATA / "categorized-tickets.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        out.to_excel(xl, sheet_name="Tickets", index=False)
        per_cat.to_excel(xl, sheet_name="Accuracy", index=False)
        review.to_excel(xl, sheet_name="Needs Review", index=False)
        wrong.to_excel(xl, sheet_name="Misclassified", index=False)
        style(xl.book)

    print(f"Scored {n} tickets against the Answer Key\n")
    print(f"  L1 category    : {overall['L1']}%")
    print(f"  L2 subcategory : {overall['L2']}%")
    print(f"  L3 issue type  : {overall['L3']}%")
    print(f"\n  Uncategorized  : {accuracy['uncategorized']}")
    print(f"  Needs review   : {accuracy['needs_review_tickets']} tickets "
          f"(Medium or Low confidence)")
    print(f"\nAccuracy by confidence band (L2):\n{by_conf.to_string(index=False)}")
    print(f"\nAccuracy by true category:\n{per_cat.to_string(index=False)}")
    if len(wrong):
        print(f"\n{len(wrong)} distinct descriptions landed in the wrong subcategory:")
        for _, r in wrong.iterrows():
            print(f"  [{r.Confidence:6}] {r['Subcategory (L2)']:24} "
                  f"should be {r['True Subcategory (L2)']:24} | {r.Description[:60]}")
    print(f"\nWrote {path}")


def style(book):
    fill = PatternFill("solid", fgColor="1F3B57")
    font = Font(color="FFFFFF", bold=True, size=10)
    widths = {"Description": 62, "Short Description": 30, "Reason": 52,
              "Ticket ID": 12, "Opened Date": 17, "Resolved Date": 17,
              "Category (L1)": 24, "Subcategory (L2)": 26, "Issue Type (L3)": 30,
              "True Subcategory (L2)": 26, "Assignment Group": 20, "Requester": 18}
    for ws in book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill, cell.font = fill, font
            cell.alignment = Alignment(vertical="center")
        for i, cell in enumerate(ws[1], start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(
                cell.value, max(14, min(30, len(str(cell.value)) + 4)))
        ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    main()
