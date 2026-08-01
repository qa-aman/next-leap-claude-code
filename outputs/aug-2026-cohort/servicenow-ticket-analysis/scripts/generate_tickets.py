"""
Builds the sample ServiceNow ticket dataset.

Output: data/servicenow-tickets-jan-mar-2026.xlsx with three sheets
  1. Tickets    - 450 tickets, category columns deliberately BLANK
  2. Answer Key - the true L1/L2/L3 per ticket, used only to measure categorizer accuracy
  3. Taxonomy   - the fixed 3-tier reference

Run: python3 generate_tickets.py
Deterministic (fixed seed), so re-running gives the identical file.
"""

import random
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from taxonomy import TAXONOMY, leaves

SEED = 20260317
random.seed(SEED)

OUT = Path(__file__).resolve().parent.parent / "data"
OUT.mkdir(parents=True, exist_ok=True)

MONTHS = [("2026-01", 140), ("2026-02", 150), ("2026-03", 160)]

# Relative share of total volume, per L2 subcategory. Sums to 100.
WEIGHTS = {
    "Password": 18, "Account Provisioning": 5, "MFA": 4, "Group Membership": 3,
    "Laptop": 7, "Peripherals": 5, "Printer": 3, "Mobile": 2,
    "Office Productivity Suite": 5, "ERP System": 5, "CRM System": 3,
    "Licensing": 4, "Browser and Plugins": 3,
    "VPN": 7, "WiFi": 4, "LAN and Cabling": 1.5, "Bandwidth and Performance": 1.5,
    "Mail Client": 4, "Meetings and Calls": 4, "Document Collaboration": 2,
    "Distribution Lists": 2,
    "Phishing and Fraud": 3, "Malware": 1.5, "Data Access Control": 1.5,
    "Audit and Review": 1,
}

# Month-on-month multiplier. Password climbs, VPN eases off after the client upgrade.
TREND = {
    "Password": {"2026-01": 0.75, "2026-02": 1.00, "2026-03": 1.30},
    "VPN": {"2026-01": 1.20, "2026-02": 1.00, "2026-03": 0.85},
    "Licensing": {"2026-01": 0.80, "2026-02": 1.00, "2026-03": 1.20},
}

# (median resolution hours, lognormal sigma). Small sigma = same fix every time = automatable.
RESOLUTION = {
    "Password": (0.8, 0.25), "Account Provisioning": (6.0, 0.50),
    "MFA": (1.2, 0.35), "Group Membership": (4.0, 0.60),
    "Laptop": (20.0, 0.90), "Peripherals": (8.0, 0.80),
    "Printer": (5.0, 0.90), "Mobile": (6.0, 0.70),
    "Office Productivity Suite": (3.0, 0.80), "ERP System": (14.0, 1.00),
    "CRM System": (10.0, 1.00), "Licensing": (18.0, 0.70),
    "Browser and Plugins": (3.0, 0.80),
    "VPN": (4.0, 0.90), "WiFi": (9.0, 1.00),
    "LAN and Cabling": (12.0, 0.80), "Bandwidth and Performance": (16.0, 1.00),
    "Mail Client": (2.5, 0.60), "Meetings and Calls": (4.0, 0.90),
    "Document Collaboration": (5.0, 0.80), "Distribution Lists": (1.5, 0.30),
    "Phishing and Fraud": (2.0, 0.50), "Malware": (8.0, 1.00),
    "Data Access Control": (20.0, 0.80), "Audit and Review": (30.0, 0.70),
}

# Priority mix as weights over P1..P4.
PRIORITY = {
    "Password": (0, 8, 55, 37), "Account Provisioning": (0, 15, 60, 25),
    "MFA": (0, 20, 55, 25), "Group Membership": (0, 10, 55, 35),
    "Laptop": (2, 25, 55, 18), "Peripherals": (0, 10, 55, 35),
    "Printer": (0, 8, 50, 42), "Mobile": (0, 10, 55, 35),
    "Office Productivity Suite": (0, 20, 55, 25), "ERP System": (8, 40, 42, 10),
    "CRM System": (3, 30, 52, 15), "Licensing": (0, 10, 45, 45),
    "Browser and Plugins": (0, 8, 55, 37),
    "VPN": (3, 35, 48, 14), "WiFi": (2, 25, 55, 18),
    "LAN and Cabling": (0, 20, 60, 20), "Bandwidth and Performance": (5, 40, 45, 10),
    "Mail Client": (0, 22, 55, 23), "Meetings and Calls": (2, 30, 53, 15),
    "Document Collaboration": (0, 15, 60, 25), "Distribution Lists": (0, 5, 45, 50),
    "Phishing and Fraud": (25, 50, 20, 5), "Malware": (40, 45, 13, 2),
    "Data Access Control": (0, 15, 55, 30), "Audit and Review": (0, 20, 55, 25),
}

SLA_TARGET = {"P1 - Critical": 4, "P2 - High": 8, "P3 - Moderate": 16, "P4 - Low": 40}
PRIORITY_LABELS = ["P1 - Critical", "P2 - High", "P3 - Moderate", "P4 - Low"]

REOPEN_RATE = {
    "VPN": 0.14, "WiFi": 0.12, "Laptop": 0.11, "Bandwidth and Performance": 0.13,
    "ERP System": 0.09, "CRM System": 0.08, "Printer": 0.07,
}
DEFAULT_REOPEN = 0.03

CHANNELS = (["Self-Service Portal"] * 46 + ["Email"] * 27 + ["Phone"] * 18 + ["Chat"] * 9)

DEPARTMENTS = [
    "Sales", "Finance", "Human Resources", "Operations", "Engineering",
    "Customer Success", "Marketing", "Legal", "Procurement", "Manufacturing",
]

FIRST = ["Priya", "Rahul", "Anita", "Vikram", "Sneha", "Arjun", "Meera", "Karan",
         "Divya", "Rohit", "Fatima", "Daniel", "Grace", "Miguel", "Chloe", "Tobias",
         "Ayesha", "Nikhil", "Laura", "Omar", "Ingrid", "Samuel", "Nadia", "Peter"]
LAST = ["Sharma", "Iyer", "Nair", "Bose", "Kapoor", "Menon", "Reddy", "Joshi",
        "Okafor", "Silva", "Novak", "Fischer", "Dubois", "Rossi", "Andersen",
        "Haddad", "Kowalski", "Tanaka", "Mensah", "Ferreira"]

# Filler that real users add and that carries no categorisation signal.
PREFIX = ["", "", "", "hi team, ", "hello, ", "hi, ", "team, ", "good morning, ",
          "raising this again, ", "second time reporting this, "]
SUFFIX = ["", "", "", "", " please help", " kindly look into it urgently",
          " need this today", " this is blocking my work", " thanks in advance",
          " let me know if you need anything from my side", " appreciate a quick fix"]

# Vague short descriptions. A third of tickets carry one of these instead of anything useful,
# which is exactly why the long description has to be read.
VAGUE_SHORT = ["Issue", "Not working", "Urgent help needed", "Please assist",
               "Request", "System problem", "Need support", "IT issue"]

TYPO_MAP = {"the": "teh", "and": "adn", "please": "plz", "cannot": "cannt",
            "working": "workign", "system": "sytem", "not": "nto",
            "morning": "mornign", "access": "acces", "again": "agian"}


def add_noise(text):
    """Make the description look typed in a hurry."""
    out = random.choice(PREFIX) + text + random.choice(SUFFIX)
    if random.random() < 0.35:
        words = out.split()
        for i, w in enumerate(words):
            if w in TYPO_MAP and random.random() < 0.5:
                words[i] = TYPO_MAP[w]
        out = " ".join(words)
    if random.random() < 0.12:
        out = out.upper()
    return out


def short_desc(text):
    if random.random() < 0.32:
        return random.choice(VAGUE_SHORT)
    words = text.split()
    return " ".join(words[:6]).rstrip(",.").capitalize()


def lognormal_hours(median, sigma):
    return round(max(0.15, median * random.lognormvariate(0, sigma)), 2)


def business_datetime(month, rng):
    y, m = int(month[:4]), int(month[5:])
    last = {1: 31, 2: 28, 3: 31}[m]
    while True:
        d = rng.randint(1, last)
        dt = datetime(y, m, d)
        if dt.weekday() < 5:
            break
    hour = rng.choices(range(8, 19), weights=[6, 12, 14, 11, 9, 7, 9, 10, 8, 5, 3])[0]
    return dt.replace(hour=hour, minute=rng.randint(0, 59))


def build():
    leaf_list = list(leaves())
    by_l2 = {}
    for l1, l2, l3, phrasings, group in leaf_list:
        by_l2.setdefault(l2, {"l1": l1, "group": group, "issues": []})
        by_l2[l2]["issues"].append((l3, phrasings))

    rows, key_rows = [], []
    counter = 1045200

    for month, target in MONTHS:
        weighted = []
        for l2, w in WEIGHTS.items():
            weighted.append((l2, w * TREND.get(l2, {}).get(month, 1.0)))
        total_w = sum(w for _, w in weighted)
        picks = random.choices(
            [l2 for l2, _ in weighted],
            weights=[w / total_w for _, w in weighted],
            k=target,
        )

        for l2 in picks:
            meta = by_l2[l2]
            l1, group = meta["l1"], meta["group"]
            l3, phrasings = random.choice(meta["issues"])
            raw = random.choice(phrasings)

            counter += random.randint(1, 4)
            tid = f"INC{counter:07d}"

            opened = business_datetime(month, random)
            priority = random.choices(PRIORITY_LABELS, weights=PRIORITY[l2])[0]
            med, sig = RESOLUTION[l2]
            hours = lognormal_hours(med, sig)

            # A slice of March is still open, as it would be mid-month.
            open_state = month == "2026-03" and random.random() < 0.09
            if open_state:
                state = random.choice(["In Progress", "On Hold", "New"])
                resolved, hours_out, sla_met = "", "", ""
            else:
                state = random.choices(["Resolved", "Closed"], weights=[35, 65])[0]
                resolved = opened + timedelta(hours=hours)
                hours_out = hours
                sla_met = "Yes" if hours <= SLA_TARGET[priority] else "No"

            reopened = "Yes" if (not open_state and random.random() <
                                 REOPEN_RATE.get(l2, DEFAULT_REOPEN)) else "No"

            description = add_noise(raw)
            rows.append({
                "Ticket ID": tid,
                "Opened Date": opened.strftime("%d-%m-%Y %H:%M"),
                "Month": month,
                "Requester": f"{random.choice(FIRST)} {random.choice(LAST)}",
                "Department": random.choice(DEPARTMENTS),
                "Channel": random.choice(CHANNELS),
                "Short Description": short_desc(raw),
                "Description": description,
                "Priority": priority,
                "Assignment Group": group,
                "State": state,
                "Resolved Date": resolved.strftime("%d-%m-%Y %H:%M") if resolved else "",
                "Resolution Hours": hours_out,
                "SLA Target Hours": SLA_TARGET[priority],
                "SLA Met": sla_met,
                "Reopened": reopened,
                "Category (L1)": "",
                "Subcategory (L2)": "",
                "Issue Type (L3)": "",
            })
            key_rows.append({
                "Ticket ID": tid, "Month": month,
                "True Category (L1)": l1,
                "True Subcategory (L2)": l2,
                "True Issue Type (L3)": l3,
            })

    tickets = pd.DataFrame(rows).sort_values("Opened Date", key=lambda s: pd.to_datetime(
        s, format="%d-%m-%Y %H:%M")).reset_index(drop=True)
    key = pd.DataFrame(key_rows).set_index("Ticket ID").loc[tickets["Ticket ID"]].reset_index()

    tax_rows = [
        {"Category (L1)": l1, "Subcategory (L2)": l2, "Issue Type (L3)": l3,
         "Assignment Group": g, "Example Phrasing": p[0]}
        for l1, l2, l3, p, g in leaf_list
    ]
    tax = pd.DataFrame(tax_rows)

    path = OUT / "servicenow-tickets-jan-mar-2026.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        tickets.to_excel(xl, sheet_name="Tickets", index=False)
        key.to_excel(xl, sheet_name="Answer Key", index=False)
        tax.to_excel(xl, sheet_name="Taxonomy", index=False)
        style(xl.book)

    print(f"Wrote {path}")
    print(f"  Tickets: {len(tickets)}")
    for month, _ in MONTHS:
        print(f"    {month}: {(tickets['Month'] == month).sum()}")
    print("\nTrue L1 mix:")
    print(key["True Category (L1)"].value_counts().to_string())
    return path


def style(book):
    head_fill = PatternFill("solid", fgColor="1F3B57")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    widths = {"Description": 62, "Short Description": 30, "Ticket ID": 12,
              "Opened Date": 17, "Resolved Date": 17, "Requester": 18,
              "Assignment Group": 20, "Priority": 15, "Department": 18,
              "Channel": 19, "Example Phrasing": 55}
    for ws in book.worksheets:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(vertical="center")
        for i, cell in enumerate(ws[1], start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(
                cell.value, max(14, min(30, len(str(cell.value)) + 4)))
        ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    build()
